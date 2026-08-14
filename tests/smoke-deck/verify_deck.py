#!/usr/bin/env python3
"""Verify a codex-ppt smoke-deck run.

Two tiers:
  BASELINE  must pass against the unmodified skill. The exit code reflects this tier only.
  TARGET    describes behavior the skill does not have yet; reported as a progress counter.

See README.md for the run protocol.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

SLIDE_COUNT = 3
DECK_NAME = "coffee-extraction"

TERMINAL_STATUSES = {"recorded", "accepted", "blocked"}
IN_FLIGHT_STATUSES = {"pending", "dispatched"}

HOST_ACCEPTABLE = {
    "codex": {"builtin-image-gen"},
    # The unmodified skill's own subagent handoff template (prompts/slide-worker.md) only ever
    # names two generic backend families -- "built-in image tool" or "scripts/image_gen.py" --
    # so a doc-compliant subagent reporting "built-in image tool" on these hosts must still pass.
    # The specific-tool names (codex-gpt-image, GenerateImage) are also accepted when present;
    # rewarding that specificity is what TARGET T6 measures.
    # Cursor GenerateImage is 3:2-only; a 16:9/4:3 GenerateImage run still fails the image-ratio
    # baseline when pixels are 1536x1024. This matrix is host-label acceptance, not a ratio gate.
    "claude-code": {"builtin-image-gen", "codex-gpt-image"},
    "cursor": {"builtin-image-gen", "generateimage", "codex-gpt-image"},
}

# Tokens that prove a backend label names a specific tool rather than a generic phrase (T6).
# Note: OpenClaw `image_generate` normalizes to `imagegenerate`, which contains `imagegen` —
# `is_specific_backend_label()` excludes that case so T6 does not reward the dropped backend.
SPECIFIC_BACKEND_TOKENS = ("imagegen", "generateimage", "codexgptimage")

SPEECH_HEADING_RE = re.compile(
    r"^#{2,4}\s*(?:Slide\s*(\d+)|第\s*(\d+)\s*页)", re.IGNORECASE | re.MULTILINE
)

PASS, FAIL, SKIP, INFO = "[PASS]", "[FAIL]", "[SKIP]", "[INFO]"

_SCRIPTS = Path(__file__).resolve().parents[2] / "skills" / "codex-ppt" / "scripts"
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))
from slide_image_ratio import RATIO_BY_LABEL, RATIO_RELATIVE_TOLERANCE, is_ratio

MEDAVIS_LOGO_TOP_EMU = 95250
MEDAVIS_LOGO_WIDTH_EMU = 1362456
PPT_TITLE_TOKEN = "<PPT-Title>"


# --------------------------------------------------------------------------- helpers


def normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def backend_family(value: str) -> str:
    """Classify a free-form backend label. Order matters: CLI fallback is checked first
    because 'scripts/image_gen.py' also contains the 'imagegen' token.

    OpenClaw `image_generate` → normalized `imagegenerate` would otherwise match `imagegen`
    and look like Codex builtin. Align with fork skill: treat it as unknown/unsupported.
    """
    n = normalize(value)
    if "imagegenpy" in n or "scriptsimagen" in n or "scriptsimagegen" in n:
        return "cli-api-fallback"
    if "codexgptimage" in n:
        return "codex-gpt-image"
    if "generateimage" in n:
        return "generateimage"
    if "imagegenerate" in n:
        return "unknown"
    if "imagegen" in n or "builtin" in n:
        return "builtin-image-gen"
    return "unknown"


def is_specific_backend_label(value: str) -> bool:
    """T6 helper: specific tool name, not generic built-in phrase, not OpenClaw."""
    n = normalize(value)
    if "imagegenerate" in n and "generateimage" not in n:
        return False
    if n in {"builtinimagetool", "builtinimage", "builtin"}:
        return False
    return any(token in n for token in SPECIFIC_BACKEND_TOKENS)


def method_backend_label(method: dict | None) -> str | None:
    """Mirror prepare_slide_prompts.py's _method_backend_label(): the backend name may be
    recorded under any of these keys, not just 'backend_used'."""
    if not method:
        return None
    for key in ("backend_used", "selected_backend", "backend", "tool_name"):
        value = method.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def cjk_count(text: str) -> int:
    return sum(1 for ch in text if "一" <= ch <= "鿿" or "㐀" <= ch <= "䶿")


def latin_count(text: str) -> int:
    return sum(1 for ch in text if ch.isascii() and ch.isalpha())


def cjk_ratio(text: str) -> float:
    cjk, latin = cjk_count(text), latin_count(text)
    total = cjk + latin
    return cjk / total if total else 0.0


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def ratio_matches(width: float, height: float, expected: float, tol: float = RATIO_RELATIVE_TOLERANCE) -> bool:
    return is_ratio(width, height, expected, tol)


def locked_aspect_ratio(deck: Path) -> str:
    for name in ("deck_spec.json", "slide_jobs.json"):
        path = deck / name
        if not path.exists():
            continue
        try:
            data = read_json(path)
        except (OSError, json.JSONDecodeError, TypeError):
            continue
        if isinstance(data, dict):
            label = str(data.get("aspect_ratio") or "").strip()
            if label in RATIO_BY_LABEL:
                return label
    return "16:9"


def runtime_python() -> Path:
    home = Path(os.environ.get("CODEX_PPT_HOME", "~/.codex-ppt-skill")).expanduser()
    if os.name == "nt":
        return home / ".venv" / "Scripts" / "python.exe"
    return home / ".venv" / "bin" / "python"


class Report:
    def __init__(self) -> None:
        self.baseline: list[tuple[bool, str, str]] = []
        self.target: list[tuple[bool | None, str, str]] = []
        self.info: list[str] = []

    def base(self, ok: bool, name: str, detail: str = "") -> None:
        self.baseline.append((ok, name, detail))

    def targ(self, ok: bool | None, name: str, detail: str = "") -> None:
        self.target.append((ok, name, detail))

    def note(self, text: str) -> None:
        self.info.append(text)

    @property
    def baseline_failed(self) -> int:
        return sum(1 for ok, _, _ in self.baseline if not ok)

    def render(self) -> None:
        print("=" * 72)
        print("BASELINE  (must pass against the unmodified skill)")
        print("=" * 72)
        for ok, name, detail in self.baseline:
            print(f"{PASS if ok else FAIL} {name}" + (f" -- {detail}" if detail else ""))

        satisfied = sum(1 for ok, _, _ in self.target if ok is True)
        scored = sum(1 for ok, _, _ in self.target if ok is not None)
        print()
        print("=" * 72)
        print(f"TARGET  {satisfied}/{scored} satisfied  (roadmap; does not affect exit code)")
        print("=" * 72)
        for ok, name, detail in self.target:
            marker = SKIP if ok is None else (PASS if ok else FAIL)
            print(f"{marker} {name}" + (f" -- {detail}" if detail else ""))

        if self.info:
            print()
            print("=" * 72)
            print("INFO")
            print("=" * 72)
            for line in self.info:
                print(f"{INFO} {line}")

        print()
        failed = self.baseline_failed
        if failed:
            print(f"RESULT: BASELINE FAILED ({failed} check(s)). Something regressed.")
        else:
            print(f"RESULT: BASELINE PASSED. TARGET progress {satisfied}/{scored}.")


# --------------------------------------------------------------------------- preflight


def cmd_preflight() -> int:
    print("=" * 72)
    print("PREFLIGHT")
    print("=" * 72)

    ok = True
    py = runtime_python()
    if py.exists():
        print(f"{PASS} shared runtime interpreter: {py}")
        probe = subprocess.run(
            [str(py), "-c", "import pptx, PIL, filelock"],
            capture_output=True,
            text=True,
        )
        if probe.returncode == 0:
            print(f"{PASS} runtime deps importable (pptx, PIL, filelock)")
        else:
            ok = False
            missing = probe.stderr.strip().splitlines()[-1] if probe.stderr.strip() else "unknown"
            print(f"{FAIL} runtime deps NOT importable -- {missing}")
    else:
        ok = False
        print(f"{FAIL} shared runtime interpreter missing: {py}")

    try:
        import pptx  # noqa: F401
        import PIL  # noqa: F401

        print(f"{PASS} verifier deps importable in this interpreter ({sys.executable})")
    except ImportError as exc:
        ok = False
        print(f"{FAIL} verifier deps missing in this interpreter -- {exc}")
        print(f"       run the verifier with: {py}")

    if not ok:
        print()
        print("To create or repair the shared runtime, run:")
        print(f"  python3 <skill_root>/scripts/codex_ppt_runtime.py bootstrap")
        print()
        print("Note: on Windows the interpreter lives at .venv/Scripts/python.exe;")
        print("skill docs should use {runtime_python}, not a Unix-only path.")
    return 0 if ok else 1


# --------------------------------------------------------------------------- checks


def check_artifacts(deck: Path, r: Report) -> None:
    required = [
        "outline.md",
        "deck_spec.json",
        "speech.md",
        "slide_jobs.json",
        "slide_run_state.json",
        f"{DECK_NAME}.pptx",
    ] + [f"prompts/slide_{i:02d}.json" for i in range(1, SLIDE_COUNT + 1)]

    missing = [name for name in required if not (deck / name).exists()]
    r.base(not missing, "artifacts present", f"missing: {', '.join(missing)}" if missing else "")

    image_dir = deck / "origin_image"
    if not image_dir.is_dir():
        r.base(False, "origin_image/ holds exactly the final slides", "origin_image/ missing")
        return
    expected = {f"slide_{i:02d}.png" for i in range(1, SLIDE_COUNT + 1)}
    actual = {p.name for p in image_dir.iterdir() if p.is_file()}
    extra, absent = sorted(actual - expected), sorted(expected - actual)
    detail = "; ".join(
        part
        for part in (
            f"unexpected: {', '.join(extra)}" if extra else "",
            f"missing: {', '.join(absent)}" if absent else "",
        )
        if part
    )
    r.base(not extra and not absent, "origin_image/ holds exactly the final slides", detail)


def check_state_and_provenance(deck: Path, host: str, allow_cli: bool, r: Report) -> None:
    jobs_path = deck / "slide_jobs.json"
    if not jobs_path.exists():
        r.base(False, "slide state", "slide_jobs.json missing")
        return
    jobs = read_json(jobs_path)
    slides = jobs.get("slides", [])

    r.base(len(slides) == SLIDE_COUNT, "slide count is 3", f"found {len(slides)}")

    statuses = [s.get("status") for s in slides]
    stuck = [s for s in statuses if s in IN_FLIGHT_STATUSES]
    r.base(
        all(s in TERMINAL_STATUSES for s in statuses),
        "all slide statuses terminal",
        f"still in flight: {', '.join(stuck)}" if stuck else "",
    )

    blocked = [s for s in slides if s.get("status") == "blocked"]
    r.base(
        not blocked,
        "no slide blocked",
        "; ".join(
            f"{s.get('slide_id')}: {(s.get('blocker') or {}).get('reason', '?')}" for s in blocked
        ),
    )

    recorded = sum(1 for s in statuses if s == "recorded")
    r.base(
        recorded >= 2,
        "subagent path exercised (>=2 slides recorded)",
        f"recorded={recorded}",
    )

    families: set[str] = set()
    bad_backend: list[str] = []
    unmatched: list[str] = []
    acceptable = HOST_ACCEPTABLE[host]

    sample_backend = method_backend_label(jobs.get("sample_generation_method"))
    if sample_backend:
        sample_family = backend_family(sample_backend)
        families.add(sample_family)
        sample_allowed = sample_family in acceptable or (sample_family == "cli-api-fallback" and allow_cli)
        if not sample_allowed:
            bad_backend.append(f"sample={sample_backend!r} -> {sample_family}")
    else:
        bad_backend.append("sample_generation_method missing a backend label")

    for slide in slides:
        result = slide.get("result") or {}
        backend = result.get("backend_used")
        if not backend:
            continue  # accepted sample may carry no worker result
        family = backend_family(backend)
        families.add(family)
        allowed = family in acceptable or (family == "cli-api-fallback" and allow_cli)
        if not allowed:
            bad_backend.append(f"{slide.get('slide_id')}={backend!r} -> {family}")
        if result.get("sample_generation_method_matched") is False:
            unmatched.append(str(slide.get("slide_id")))

    r.base(
        not bad_backend,
        f"backend acceptable for host '{host}'",
        "; ".join(bad_backend) if bad_backend else f"families={sorted(families) or ['-']}",
    )
    r.base(
        len(families) <= 1,
        "one backend family across the deck",
        f"families={sorted(families)}" if len(families) > 1 else "",
    )
    r.base(
        not unmatched,
        "sample_generation_method matched",
        f"unmatched: {', '.join(unmatched)}" if unmatched else "",
    )

    if not allow_cli and "cli-api-fallback" in families:
        r.note("CLI/API fallback was used without --allow-cli-fallback: backend detection likely failed.")

    # --- targets fed by the same file
    accepted = sum(1 for s in statuses if s == "accepted")
    r.targ(accepted == 1, "T4 exactly 1 slide accepted (sample reused, not regenerated)",
           f"accepted={accepted}")

    selected = jobs.get("selected_backend") or ""
    r.targ(
        is_specific_backend_label(selected),
        "T6 selected_backend names a specific tool",
        f"selected_backend={selected!r}",
    )


def check_deck_spec(deck: Path, r: Report) -> None:
    spec_path = deck / "deck_spec.json"
    if not spec_path.exists():
        r.base(False, "deck_spec structure", "deck_spec.json missing")
        return
    spec = read_json(spec_path)
    slides = spec.get("slides", [])
    by_number = {}
    for index, slide in enumerate(slides, start=1):
        try:
            by_number[int(slide.get("number", index))] = slide
        except (TypeError, ValueError):
            by_number[index] = slide

    roles = {str(s.get("role") or "").strip().lower() for s in slides}
    roles.discard("")
    r.base(len(roles) >= 2, "layout roles vary (>=2 distinct)", f"roles={sorted(roles) or ['-']}")

    def slide_text(slide: dict) -> str:
        points = slide.get("key_points") or []
        if isinstance(points, str):
            points = [points]
        return " ".join([str(slide.get("title") or "")] + [str(p) for p in points])

    expectations = {
        1: ("Chinese", lambda t: cjk_ratio(t) >= 0.30),
        2: ("English", lambda t: cjk_ratio(t) <= 0.05),
        3: ("mixed", lambda t: cjk_count(t) > 0 and latin_count(t) > 0),
    }
    problems = []
    for number, (label, predicate) in expectations.items():
        slide = by_number.get(number)
        if slide is None:
            problems.append(f"slide {number} missing")
            continue
        text = slide_text(slide)
        if not predicate(text):
            problems.append(f"slide {number} not {label} (cjk_ratio={cjk_ratio(text):.2f})")
    r.base(not problems, "per-slide language contract honored", "; ".join(problems))

    # --- target
    with_language = [s for s in slides if str(s.get("language") or "").strip()]
    r.targ(
        len(with_language) == len(slides) and bool(slides),
        "T3 per-slide language field present in deck_spec",
        f"{len(with_language)}/{len(slides)} slides carry it",
    )


def check_speech(deck: Path, r: Report) -> None:
    speech = deck / "speech.md"
    if not speech.exists():
        r.base(False, "speech.md slide headings", "speech.md missing")
        return
    found = {int(a or b) for a, b in SPEECH_HEADING_RE.findall(speech.read_text(encoding="utf-8"))}
    expected = set(range(1, SLIDE_COUNT + 1))
    r.base(
        found >= expected,
        "speech.md has parseable headings for all slides",
        f"found slides {sorted(found) or ['-']}",
    )


def check_pptx(deck: Path, r: Report) -> None:
    path = deck / f"{DECK_NAME}.pptx"
    if not path.exists():
        r.base(False, "pptx structure", f"{DECK_NAME}.pptx missing")
        return
    try:
        from pptx import Presentation
    except ImportError:
        r.base(False, "pptx structure", "python-pptx not importable; run via the shared runtime python")
        return

    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        picture_type = MSO_SHAPE_TYPE.PICTURE
    except ImportError:  # pragma: no cover
        picture_type = 13

    prs = Presentation(str(path))
    slides = list(prs.slides)
    r.base(len(slides) == SLIDE_COUNT, "pptx has 3 slides", f"found {len(slides)}")

    label = locked_aspect_ratio(deck)
    expected = RATIO_BY_LABEL[label]
    ok = ratio_matches(float(prs.slide_width), float(prs.slide_height), expected)
    r.base(ok, f"pptx page size is {label}", f"ratio={float(prs.slide_width) / float(prs.slide_height):.4f}")

    no_picture, no_notes = [], []
    for index, slide in enumerate(slides, start=1):
        if not any(shape.shape_type == picture_type for shape in slide.shapes):
            no_picture.append(str(index))
        text = ""
        if slide.has_notes_slide:
            text = (slide.notes_slide.notes_text_frame.text or "").strip()
        if not text:
            no_notes.append(str(index))

    r.base(not no_picture, "every pptx slide carries a picture",
           f"missing on slide(s): {', '.join(no_picture)}" if no_picture else "")
    r.base(not no_notes, "every pptx slide carries speaker notes",
           f"empty on slide(s): {', '.join(no_notes)}" if no_notes else "")


# Medavis wrap: landing + N content + ending. Overlay logo sits at 10px top margin.
MEDAVIS_PPTX_NAME = f"{DECK_NAME}-medavis.pptx"


def check_medavis_pptx(deck: Path, r: Report) -> None:
    """When coffee-extraction-medavis.pptx exists, validate landing + content wrap.

    Absent file = skip (FAST / none-only runs). Present file = baseline checks so
    interactive/both runs do not silently ship a broken medavis deck.
    """
    path = deck / MEDAVIS_PPTX_NAME
    if not path.exists():
        r.note(f"{MEDAVIS_PPTX_NAME} not present — skipping medavis pptx checks")
        return

    try:
        from pptx import Presentation
    except ImportError:
        r.base(False, "medavis pptx structure", "python-pptx not importable; run via the shared runtime python")
        return

    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        picture_type = MSO_SHAPE_TYPE.PICTURE
    except ImportError:  # pragma: no cover
        picture_type = 13

    expected_slides = 1 + SLIDE_COUNT + 1
    prs = Presentation(str(path))
    slides = list(prs.slides)
    r.base(
        len(slides) == expected_slides,
        f"medavis pptx has {expected_slides} slides (1 landing + {SLIDE_COUNT} content + 1 ending)",
        f"found {len(slides)}",
    )

    label = locked_aspect_ratio(deck)
    expected = RATIO_BY_LABEL[label]
    ok = ratio_matches(float(prs.slide_width), float(prs.slide_height), expected)
    r.base(ok, f"medavis pptx page size is {label}", f"ratio={float(prs.slide_width) / float(prs.slide_height):.4f}")

    landing = slides[0] if slides else None
    if landing is not None:
        landing_text = " ".join(
            (shape.text_frame.text or "").strip()
            for shape in landing.shapes
            if getattr(shape, "has_text_frame", False)
        ).strip()
        r.base(
            bool(landing_text) and PPT_TITLE_TOKEN not in landing_text,
            "medavis landing slide has title text (not <PPT-Title>)",
            "" if landing_text and PPT_TITLE_TOKEN not in landing_text else f"text={landing_text[:80]!r}",
        )
        landing_notes = ""
        if landing.has_notes_slide:
            landing_notes = (landing.notes_slide.notes_text_frame.text or "").strip()
        r.base(
            not landing_notes,
            "medavis landing slide has empty speaker notes",
            f"notes={landing_notes[:40]!r}" if landing_notes else "",
        )

    content_slides = slides[1:1 + SLIDE_COUNT]
    bad_pic_top, no_picture, no_notes, no_logo = [], [], [], []
    for offset, slide in enumerate(content_slides):
        index = offset + 2
        pictures = [shape for shape in slide.shapes if shape.shape_type == picture_type]
        if not pictures:
            no_picture.append(str(index))
        else:
            pic = pictures[0]
            if pic.top != 0:
                bad_pic_top.append(f"{index}(top={pic.top})")
            overlay = [
                p for p in pictures[1:]
                if abs(int(p.top) - MEDAVIS_LOGO_TOP_EMU) <= 2
                and abs(int(p.width) - MEDAVIS_LOGO_WIDTH_EMU) <= 2
            ]
            if not overlay:
                no_logo.append(str(index))
        text = ""
        if slide.has_notes_slide:
            text = (slide.notes_slide.notes_text_frame.text or "").strip()
        if not text:
            no_notes.append(str(index))

    r.base(not no_picture, "every medavis content slide carries a picture",
           f"missing on slide(s): {', '.join(no_picture)}" if no_picture else "")
    r.base(
        not bad_pic_top,
        "medavis content pictures are full-page (top==0)",
        f"not at top on slide(s): {', '.join(bad_pic_top)}" if bad_pic_top else "",
    )
    r.base(
        not no_logo,
        "medavis content slides have overlay logo near top=95250",
        f"missing on slide(s): {', '.join(no_logo)}" if no_logo else "",
    )
    r.base(not no_notes, "every medavis content slide carries speaker notes",
           f"empty on slide(s): {', '.join(no_notes)}" if no_notes else "")

    if slides:
        ending = slides[-1]
        ending_notes = ""
        if ending.has_notes_slide:
            ending_notes = (ending.notes_slide.notes_text_frame.text or "").strip()
        r.base(
            not ending_notes,
            "medavis ending slide has empty speaker notes",
            f"notes={ending_notes[:40]!r}" if ending_notes else "",
        )
        ending_text = " ".join(
            (shape.text_frame.text or "").strip()
            for shape in ending.shapes
            if getattr(shape, "has_text_frame", False)
        )
        r.base(
            "Thank You" in ending_text or "Geduld" in ending_text,
            "medavis last slide is the ending page",
            "" if ending_text else "no ending text",
        )


def check_images(deck: Path, r: Report) -> None:
    """Locked-ratio mismatch is a baseline failure. T5 >=1920px stays a target."""
    try:
        from PIL import Image
    except ImportError:
        r.base(False, "slide images match locked ratio", "Pillow not importable")
        r.targ(None, "T5 all slide images >=1920px wide", "Pillow not importable")
        return

    label = locked_aspect_ratio(deck)
    expected = RATIO_BY_LABEL[label]
    widths: list[int] = []
    mismatches: list[str] = []
    for i in range(1, SLIDE_COUNT + 1):
        path = deck / "origin_image" / f"slide_{i:02d}.png"
        if not path.exists():
            continue
        with Image.open(path) as img:
            width, height = img.size
        widths.append(width)
        if not ratio_matches(width, height, expected):
            mismatches.append(f"slide_{i:02d}.png {width}x{height}")
        r.note(f"slide_{i:02d}.png  {width}x{height}  locked={label}")

    if not widths:
        r.base(False, f"slide images match locked {label} ratio", "no images found")
        r.targ(None, "T5 all slide images >=1920px wide", "no images found")
        return
    r.base(
        not mismatches,
        f"slide images match locked {label} ratio (2% relative)",
        "; ".join(mismatches) if mismatches else "",
    )
    r.targ(min(widths) >= 1920, "T5 all slide images >=1920px wide", f"min width={min(widths)}")


def check_targets(deck: Path, skill_root: Path | None, r: Report) -> None:
    # T1 -- pre-assembly review gate
    state_path = deck / "slide_run_state.json"
    reviewed = False
    if state_path.exists():
        state = read_json(state_path)
        blob = json.dumps(state, ensure_ascii=False).lower()
        reviewed = "user_reviewed" in blob or "review_approved" in blob or "userreviewed" in blob
    r.targ(reviewed, "T1 pre-assembly review recorded in run state",
           "no review marker found" if not reviewed else "")

    # T2 -- hardcoded Chinese instruction on the English slide
    job = deck / "prompts" / "slide_02.json"
    if not job.exists():
        r.targ(None, "T2 English slide prompt free of hardcoded Chinese-text rule", "slide_02.json missing")
    else:
        prompt = str(read_json(job).get("prompt") or "")
        hit = "render chinese text" in prompt.lower()
        r.targ(not hit, "T2 English slide prompt free of hardcoded Chinese-text rule",
               "prompt still contains 'Render Chinese text'" if hit else "")

    # T7 -- cross-platform docs
    if skill_root is None:
        r.targ(None, "T7 skill docs use platform-correct interpreter paths", "--skill-root not provided")
    elif os.name != "nt":
        r.targ(None, "T7 skill docs use platform-correct interpreter paths", "only scored on Windows")
    elif not skill_root.is_dir():
        r.targ(None, "T7 skill docs use platform-correct interpreter paths", f"not a directory: {skill_root}")
    else:
        hits = [
            md.relative_to(skill_root).as_posix()
            for md in skill_root.rglob("*.md")
            if ".venv/bin/python" in md.read_text(encoding="utf-8", errors="replace")
        ]
        r.targ(not hits, "T7 skill docs use platform-correct interpreter paths",
               f"{len(hits)} file(s) hardcode .venv/bin/python: {', '.join(hits[:3])}" if hits else "")


# --------------------------------------------------------------------------- main


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("deck", nargs="?", help=f"Deck directory, e.g. runs/<run-id>/{DECK_NAME}")
    parser.add_argument("--host", choices=sorted(HOST_ACCEPTABLE), help="Environment the deck was generated in.")
    parser.add_argument("--skill-root", help="Installed skill directory; enables the T7 docs check.")
    parser.add_argument("--allow-cli-fallback", action="store_true",
                        help="Treat scripts/image_gen.py as an acceptable backend for this run.")
    parser.add_argument("--preflight", action="store_true", help="Check the runtime and exit.")
    args = parser.parse_args()

    if args.preflight:
        return cmd_preflight()

    if not args.deck or not args.host:
        parser.error("deck directory and --host are required (or pass --preflight)")

    deck = Path(args.deck).expanduser().resolve()
    if not deck.is_dir():
        print(f"{FAIL} deck directory not found: {deck}")
        return 1

    skill_root = Path(args.skill_root).expanduser().resolve() if args.skill_root else None

    print(f"deck: {deck}")
    print(f"host: {args.host}")
    print()

    r = Report()
    check_artifacts(deck, r)
    check_state_and_provenance(deck, args.host, args.allow_cli_fallback, r)
    check_deck_spec(deck, r)
    check_speech(deck, r)
    check_pptx(deck, r)
    check_medavis_pptx(deck, r)
    check_images(deck, r)
    check_targets(deck, skill_root, r)
    r.render()

    return 1 if r.baseline_failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
