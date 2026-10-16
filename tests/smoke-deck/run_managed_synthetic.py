#!/usr/bin/env python3
"""Offline managed state/assembly smoke. All PNGs and decisions are TEST fixtures.

No provider, live image model, real content approval, pixel QA or rehearsal is
represented by this run. The copied synthetic result is not a source-fidelity test.
"""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import sys
from zipfile import ZipFile
from xml.etree import ElementTree

ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT/"skills/codex-ppt/scripts"),str(ROOT/"skills/codex-ppt/scripts/tests")]
from planning_fixtures import (managed_spec,compile_fixture,dispatch_fixture,result_fixture,
                              qa_fixture,run_cli,fixture_image)
from planning_contract import require_managed_ready
from pptx import Presentation
from PIL import Image


def checked(result):
    if result.returncode:
        raise AssertionError(result.stdout+result.stderr)
    return result


def run(deck: Path) -> dict:
    deck=deck.resolve()
    deck.mkdir(parents=True,exist_ok=False)
    source=deck/"assets/TEST-source.png"
    fixture_image(source,"blue")
    spec=managed_spec()
    spec["deck_name"]=deck.name
    spec["planning"]["brief"].update(output_profile="both",landing_title="SYNTHETIC TEST")
    spec["slides"][0]["required_images"]=[{"path":str(source),"role":"SYNTHETIC source binding",
                                          "fidelity":"Test only; no model fidelity claim"}]
    compile_fixture(deck,spec)
    checked(dispatch_fixture(deck))
    checked(result_fixture(deck,color="red"))
    for sid in ("slide_01","slide_02"):
        checked(qa_fixture(deck,sid))
    checked(run_cli("record_slide_qa.py",deck,"--deck-pass","--evidence","SYNTHETIC TEST full-sequence binding"))
    checked(run_cli("record_user_review.py",deck,"--all","--approve","--evidence","SYNTHETIC TEST current image decision"))
    require_managed_ready(deck,"both","16:9","SYNTHETIC TEST")
    checked(run_cli("assemble_ppt.py",deck.parent,deck.name+".pptx","--template","both",
                    "--landing-title","SYNTHETIC TEST","--aspect-ratio","16:9"))
    outputs=[]
    for suffix,expected,offset in (("",2,0),("-medavis",4,1)):
        path=deck/(deck.name+suffix+".pptx")
        prs=Presentation(path)
        assert len(prs.slides)==expected
        embedded=[]
        for i,slide in enumerate(spec["slides"]):
            text=prs.slides[i+offset].notes_slide.notes_text_frame.text
            assert text.strip()==slide["speaker_notes"]["text"].strip(),text
            assert "\n\n" in text
            embedded.append(text)
        if offset:
            assert not prs.slides[-1].notes_slide.notes_text_frame.text.strip()
        with ZipFile(path) as archive:
            mapped=[]
            for i in range(2):
                note_part=str(prs.slides[i+offset].notes_slide.part.partname).lstrip("/")
                xml=ElementTree.fromstring(archive.read(note_part))
                strings=[n.text or "" for n in xml.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t")]
                assert any("everyone" in t or "model" in t for t in strings)
                mapped.append({"slide_number":i+offset+1,"part":note_part})
        outputs.append({"path":str(path),"slide_count":expected,"notes_offset":offset,
            "content_note_xml_mapping":mapped,"embedded_notes":embedded,
            "sha256":hashlib.sha256(path.read_bytes()).hexdigest()})
    ratios=[]
    for path in sorted((deck/"origin_image").glob("*.png")):
        with Image.open(path) as image:
            assert image.size==(1600,900)
            ratios.append({"path":str(path),"width":image.width,"height":image.height})
    report={"fixture":True,"deck":str(deck),"source_sha256":hashlib.sha256(source.read_bytes()).hexdigest(),
            "outputs":outputs,"images":ratios,"managed_ready":True,
            "limitations":["No live image generation","No real human approval",
                "QA and source-fidelity are synthetic state evidence","No timed speech or rendered lettering assessment"]}
    (deck/"TEST-inspection.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return report


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("deck",help="New ignored scratch deck directory")
    result=run(Path(parser.parse_args().deck))
    print(json.dumps(result,ensure_ascii=False))
