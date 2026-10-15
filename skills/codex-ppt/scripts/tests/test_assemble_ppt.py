#!/usr/bin/env python3
"""Unit tests for assemble_ppt stdout/stderr encoding safety and medavis template mode."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1]
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import assemble_ppt


def _run_main_with_gbk_stdout(argv: list[str]) -> tuple[int, str]:
    """Run assemble_ppt.main() with sys.argv set to argv and sys.stdout swapped
    for a stream whose *original* encoding is GBK (simulating a Chinese-locale
    Windows console), then return (exit_code, captured_output).

    The fix reconfigures the stream to UTF-8 before printing, so the bytes
    actually written are UTF-8 -- the caller must decode them as UTF-8, not GBK,
    to see the real content. Decoding as GBK would silently produce mojibake
    instead of failing, and an assertion against an ASCII-only substring
    wouldn't notice either way (ASCII bytes round-trip identically under both
    encodings) -- which is exactly the gap this helper exists to avoid.
    """
    buffer = io.BytesIO()
    gbk_stdout = io.TextIOWrapper(buffer, encoding="gbk", errors="strict")
    old_argv, old_stdout = sys.argv, sys.stdout
    try:
        sys.argv = argv + ["--legacy-unmanaged"]
        sys.stdout = gbk_stdout
        with unittest.TestCase().assertRaises(SystemExit) as cm:
            assemble_ppt.main()
        exit_code = cm.exception.code
    finally:
        gbk_stdout.flush()
        sys.stdout = old_stdout
        sys.argv = old_argv
    buffer.seek(0)
    return exit_code, buffer.getvalue().decode("utf-8")


def _run_main(argv: list[str]) -> tuple[int, str]:
    buffer = io.StringIO()
    old_argv, old_stdout = sys.argv, sys.stdout
    try:
        sys.argv = argv + ["--legacy-unmanaged"]
        sys.stdout = buffer
        with unittest.TestCase().assertRaises(SystemExit) as cm:
            assemble_ppt.main()
        exit_code = cm.exception.code
    finally:
        sys.stdout = old_stdout
        sys.argv = old_argv
    return exit_code, buffer.getvalue()


def _make_deck_ratio(
    tmp: str,
    size: tuple[int, int],
    name: str = "demo-deck",
    slides: int = 2,
    with_notes: bool = False,
) -> Path:
    from PIL import Image

    deck_dir = Path(tmp) / name
    origin = deck_dir / "origin_image"
    origin.mkdir(parents=True)
    for i in range(1, slides + 1):
        Image.new("RGB", size, color=(i * 40, 80, 120)).save(origin / f"slide_{i:02d}.png")
    if with_notes:
        (deck_dir / "speech.md").write_text(
            "## Slide 1: One\n\nNote for slide one.\n\n## Slide 2: Two\n\nNote for slide two.\n",
            encoding="utf-8",
        )
    return deck_dir


def _make_deck(tmp: str, name: str = "demo-deck", slides: int = 2, with_notes: bool = False) -> Path:
    return _make_deck_ratio(tmp, (32, 18), name=name, slides=slides, with_notes=with_notes)


def _picture_ratio_ok(pic, expected: float, tol: float = 0.02) -> bool:
    return abs((pic.width / pic.height) / expected - 1.0) <= tol


def _template_path(filename: str) -> Path:
    return Path(assemble_ppt.skill_root_dir()) / "assets" / "templates" / filename


def _assert_template_unmutated(test: unittest.TestCase, path: Path, before) -> None:
    after = path.stat()
    test.assertEqual(before.st_size, after.st_size)
    test.assertEqual(before.st_mtime_ns, after.st_mtime_ns)


class InitStdoutEncodingTests(unittest.TestCase):
    def test_init_does_not_crash_on_gbk_console(self) -> None:
        """Windows consoles on a Chinese locale default to the GBK code page, which
        cannot encode the U+2713 check mark that --init prints alongside Chinese
        status text. A subprocess with stdout redirected through such a console
        must not crash with UnicodeEncodeError, and the printed status text must
        still be the real Chinese content, not silently-corrupted mojibake."""
        with tempfile.TemporaryDirectory() as tmp:
            exit_code, output = _run_main_with_gbk_stdout(
                ["assemble_ppt.py", tmp, "coffee-extraction.pptx", "--init"]
            )
        self.assertEqual(exit_code, 0)
        self.assertIn("✓ PPT 项目目录已准备好", output)
        self.assertIn("origin_image", output)


class AssembleStdoutEncodingTests(unittest.TestCase):
    def test_full_assembly_does_not_crash_on_gbk_console(self) -> None:
        """The non---init assembly path prints several more Chinese/check-mark
        status lines than --init alone (scan message, per-slide "already added"
        lines, final save summary). Exercise that path for real -- not just
        --init -- under the same simulated GBK console."""
        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp:
            deck_dir = Path(tmp) / "coffee-extraction"
            origin_image_dir = deck_dir / "origin_image"
            origin_image_dir.mkdir(parents=True)
            Image.new("RGB", (16, 9), color="white").save(origin_image_dir / "slide_01.png")

            exit_code, output = _run_main_with_gbk_stdout(
                ["assemble_ppt.py", tmp, "coffee-extraction.pptx"]
            )

            self.assertEqual(exit_code, 0)
            self.assertTrue((deck_dir / "coffee-extraction.pptx").exists())
        self.assertIn("正在扫描", output)  # "正在扫描" (scanning...)
        self.assertIn("已保存", output)  # "已保存" (saved)
        self.assertIn("✓", output)


class MedavisTemplateAssemblyTests(unittest.TestCase):
    def test_none_still_writes_plain_pptx(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=2)
            exit_code, _ = _run_main(
                ["assemble_ppt.py", tmp, "demo-deck.pptx", "--template", "none"]
            )
            self.assertEqual(exit_code, 0)
            plain = Path(tmp) / "demo-deck" / "demo-deck.pptx"
            medavis = Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"
            self.assertTrue(plain.exists())
            self.assertFalse(medavis.exists())
            prs = Presentation(str(plain))
            self.assertEqual(len(prs.slides), 2)
            for idx in range(len(prs.slides)):
                slide = prs.slides[idx]
                pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
                self.assertTrue(pictures)
                pic = pictures[0]
                self.assertEqual(pic.left, 0)
                self.assertEqual(pic.top, 0)
                self.assertTrue(_picture_ratio_ok(pic, 16 / 9))

    def test_medavis_writes_suffix_file_only(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=2)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                ]
            )
            self.assertEqual(exit_code, 0)
            plain = Path(tmp) / "demo-deck" / "demo-deck.pptx"
            medavis = Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"
            self.assertFalse(plain.exists())
            self.assertTrue(medavis.exists())
            prs = Presentation(str(medavis))
            self.assertEqual(len(prs.slides), 4)  # landing + 2 content + ending

    def test_both_writes_plain_and_medavis(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "both",
                    "--landing-title",
                    "Demo Title",
                ]
            )
            self.assertEqual(exit_code, 0)
            self.assertTrue((Path(tmp) / "demo-deck" / "demo-deck.pptx").exists())
            self.assertTrue((Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx").exists())

    def test_medavis_landing_title_and_color(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"))
            landing = prs.slides[0]
            landing_text = "\n".join(
                s.text_frame.text for s in landing.shapes if s.has_text_frame
            )
            self.assertIn("Demo Title", landing_text)
            self.assertNotIn("<PPT-Title>", landing_text)
            title_shapes = [
                s for s in landing.shapes if s.has_text_frame and "Demo Title" in s.text_frame.text
            ]
            self.assertTrue(title_shapes, "landing title text missing")
            alignment = title_shapes[0].text_frame.paragraphs[0].alignment
            if alignment is not None:
                from pptx.enum.text import PP_ALIGN

                self.assertEqual(alignment, PP_ALIGN.CENTER)

    def test_medavis_content_image_placement(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=2)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"))
            for idx in range(1, len(prs.slides) - 1):
                slide = prs.slides[idx]
                pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
                self.assertGreaterEqual(len(pictures), 2, "content image + overlay logo")
                pic = pictures[0]
                self.assertEqual(pic.left, 0)
                self.assertEqual(pic.top, 0)
                self.assertTrue(_picture_ratio_ok(pic, 16 / 9))
                logo = pictures[1]
                self.assertAlmostEqual(logo.top, 95250, delta=2)
                self.assertAlmostEqual(logo.width, 1362456, delta=2)
                self.assertAlmostEqual(logo.height, 274320, delta=2)

    def test_medavis_notes_mapping(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=2, with_notes=True)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"))
            landing_notes = prs.slides[0].notes_slide.notes_text_frame.text.strip()
            self.assertEqual(landing_notes, "")
            content1 = prs.slides[1].notes_slide.notes_text_frame.text
            self.assertIn("Note for slide one", content1)
            content2 = prs.slides[2].notes_slide.notes_text_frame.text
            self.assertIn("Note for slide two", content2)
            last_notes = prs.slides[len(prs.slides) - 1].notes_slide.notes_text_frame.text.strip()
            self.assertEqual(last_notes, "")

    def test_medavis_requires_landing_title(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=1)
            exit_code, _ = _run_main(
                ["assemble_ppt.py", tmp, "demo-deck.pptx", "--template", "medavis"]
            )
            self.assertNotEqual(exit_code, 0)
            self.assertFalse((Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx").exists())

    def test_both_requires_landing_title(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=1)
            exit_code, _ = _run_main(
                ["assemble_ppt.py", tmp, "demo-deck.pptx", "--template", "both"]
            )
            self.assertNotEqual(exit_code, 0)
            self.assertFalse((Path(tmp) / "demo-deck" / "demo-deck.pptx").exists())
            self.assertFalse((Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx").exists())

    def test_medavis_uses_4_3_template(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck_ratio(tmp, (32, 24), slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                    "--aspect-ratio",
                    "4:3",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"))
            self.assertEqual(prs.slide_width, 9144000)
            self.assertEqual(prs.slide_height, 6858000)

    def test_medavis_does_not_mutate_bundled_template(self) -> None:
        templates = [
            _template_path("medavis_ppt_template_2026_en_16-9.pptx"),
            _template_path("medavis_ppt_template_2026_en_4-3.pptx"),
            _template_path("medavis_ppt_template_2026_en_3-2.pptx"),
        ]
        before = {path: path.stat() for path in templates}
        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                ]
            )
            self.assertEqual(exit_code, 0)
        for path in templates:
            _assert_template_unmutated(self, path, before[path])

    def test_none_refuses_mismatched_ratio(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            _make_deck_ratio(tmp, (1536, 1024), slides=1)
            exit_code, output = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "none",
                    "--aspect-ratio",
                    "16:9",
                ]
            )
            self.assertNotEqual(exit_code, 0)
            self.assertFalse((Path(tmp) / "demo-deck" / "demo-deck.pptx").exists())
            self.assertIn("1536", output)
            self.assertIn("1024", output)

    def test_medavis_refuses_mismatched_ratio(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            _make_deck_ratio(tmp, (1536, 1024), slides=1)
            exit_code, output = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                    "--aspect-ratio",
                    "16:9",
                ]
            )
            self.assertNotEqual(exit_code, 0)
            self.assertFalse((Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx").exists())
            self.assertIn("1536", output)
            self.assertIn("1024", output)

    def test_none_3_2_accepted(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck_ratio(tmp, (48, 32), slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "none",
                    "--aspect-ratio",
                    "3:2",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck.pptx"))
            self.assertEqual(prs.slide_width, 9144000)
            self.assertEqual(prs.slide_height, 6096000)
            pic = [
                s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
            ][0]
            self.assertEqual(pic.left, 0)
            self.assertEqual(pic.top, 0)
            self.assertTrue(_picture_ratio_ok(pic, 3 / 2))

    def test_none_4_3_accepted(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck_ratio(tmp, (32, 24), slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "none",
                    "--aspect-ratio",
                    "4:3",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck.pptx"))
            self.assertEqual(prs.slide_width, 9144000)
            self.assertEqual(prs.slide_height, 6858000)
            pic = [
                s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
            ][0]
            self.assertEqual(pic.left, 0)
            self.assertEqual(pic.top, 0)
            self.assertTrue(_picture_ratio_ok(pic, 4 / 3))

    def test_none_16_9_rejected_for_aspect_4_3(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            _make_deck_ratio(tmp, (32, 18), slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "none",
                    "--aspect-ratio",
                    "4:3",
                ]
            )
            self.assertNotEqual(exit_code, 0)
            self.assertFalse((Path(tmp) / "demo-deck" / "demo-deck.pptx").exists())

    def test_medavis_3_2_uses_matching_template(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck_ratio(tmp, (48, 32), slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                    "--aspect-ratio",
                    "3:2",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"))
            self.assertEqual(prs.slide_width, 9144000)
            self.assertEqual(prs.slide_height, 6096000)
            self.assertEqual(len(prs.slides), 3)
            last = prs.slides[2]
            last_text = "\n".join(s.text_frame.text for s in last.shapes if s.has_text_frame)
            self.assertIn("Thank You For Your Patience", last_text)
            content = prs.slides[1]
            pictures = [s for s in content.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertGreaterEqual(len(pictures), 2)
            self.assertEqual(pictures[0].top, 0)
            self.assertAlmostEqual(pictures[1].top, 95250, delta=2)

    def test_medavis_ending_stays_last(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=2)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"))
            last = prs.slides[len(prs.slides) - 1]
            last_text = "\n".join(s.text_frame.text for s in last.shapes if s.has_text_frame)
            self.assertIn("Thank You For Your Patience", last_text)
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            pictures = [s for s in last.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
            overlay = [
                pic for pic in pictures
                if abs(pic.top - 95250) <= 2 and abs(pic.width - 1362456) <= 2
            ]
            self.assertEqual(overlay, [], "ending page must not get the small-logo overlay")
            self.assertEqual(len(pictures), 2, "ending page pictures must stay as in the template")

    def test_medavis_logo_geometry(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "medavis",
                    "--landing-title",
                    "Demo Title",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck-medavis.pptx"))
            pictures = [
                s for s in prs.slides[1].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
            ]
            self.assertGreaterEqual(len(pictures), 2)
            logo = pictures[1]
            expected_left = prs.slide_width - 1362456 - 95250
            self.assertAlmostEqual(logo.left, expected_left, delta=2)
            self.assertAlmostEqual(logo.top, 95250, delta=2)
            self.assertAlmostEqual(logo.width, 1362456, delta=2)
            self.assertAlmostEqual(logo.height, 274320, delta=2)

    def test_none_16_9_picture_is_uniform_full_page(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            _make_deck(tmp, slides=1)
            exit_code, _ = _run_main(
                [
                    "assemble_ppt.py",
                    tmp,
                    "demo-deck.pptx",
                    "--template",
                    "none",
                    "--aspect-ratio",
                    "16:9",
                ]
            )
            self.assertEqual(exit_code, 0)
            prs = Presentation(str(Path(tmp) / "demo-deck" / "demo-deck.pptx"))
            pic = [
                s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
            ][0]
            self.assertEqual(pic.left, 0)
            self.assertEqual(pic.top, 0)
            self.assertEqual(pic.width, prs.slide_width)
            self.assertEqual(pic.height, prs.slide_height)
            self.assertTrue(_picture_ratio_ok(pic, 16 / 9))


if __name__ == "__main__":
    unittest.main()
