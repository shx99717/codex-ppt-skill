"""Real bulk decisions and fail-closed managed PPTX output preflight."""
import copy
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from pptx import Presentation

SCRIPTS=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(SCRIPTS))
from planning_fixtures import (managed_spec,completed_fixture,run_cli,fixture_image,write_spec,approve_fixture)
from planning_contract import build_manifest


class ManagedAssemblyTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base=Path(self.tmp.name).resolve()
        self.deck=self.base/"deck"
        self.deck.mkdir()
        self.spec=completed_fixture(self.deck)

    def review(self,*args):
        return run_cli("record_user_review.py",self.deck,*args,"--evidence","SYNTHETIC TEST image decision")

    def assemble(self,*args):
        return run_cli("assemble_ppt.py",self.base,"deck.pptx",*args)

    def test_qa_alone_cannot_assemble_and_existing_output_is_preserved(self):
        output=self.deck/"deck.pptx"
        output.write_bytes(b"Previous output")
        result=self.assemble()
        self.assertNotEqual(result.returncode,0)
        self.assertEqual(output.read_bytes(),b"Previous output")

    def test_actual_bulk_review_records_one_event_and_keeps_sample_evidence(self):
        result=self.review("--all","--approve")
        self.assertEqual(result.returncode,0,result.stderr)
        ledger=json.loads((self.deck/"plan_approvals.json").read_text(encoding="utf-8"))
        events=[e for e in ledger["events"] if "image" in e["scopes"]]
        self.assertEqual(len(events),1)
        jobs=json.loads((self.deck/"slide_jobs.json").read_text(encoding="utf-8"))
        self.assertEqual(jobs["slides"][0]["user_review"]["event_id"],events[0]["id"])
        self.assertNotIn("user_review",jobs["slides"][1])
        self.assertEqual(self.assemble().returncode,0)

    def test_all_except_clears_prior_approval_without_inventing_review(self):
        self.assertEqual(self.review("--all","--approve").returncode,0)
        result=self.review("--all","--except","slide_01","--approve")
        self.assertEqual(result.returncode,0,result.stderr)
        jobs=json.loads((self.deck/"slide_jobs.json").read_text(encoding="utf-8"))
        self.assertNotIn("user_review",jobs["slides"][0])
        self.assertNotEqual(self.assemble().returncode,0)

    def test_unknown_or_incomplete_bulk_selection_does_not_mutate(self):
        before={n:(self.deck/n).read_bytes() for n in ("slide_jobs.json","plan_approvals.json")}
        result=self.review("--slides","slide_01,slide_99","--approve")
        self.assertNotEqual(result.returncode,0)
        self.assertEqual(before,{n:(self.deck/n).read_bytes() for n in before})
        jobs=json.loads(before["slide_jobs.json"])
        jobs["slides"][1]["status"]="pending"
        write_spec(self.deck/"slide_jobs.json",jobs)
        current=(self.deck/"slide_jobs.json").read_bytes()
        self.assertNotEqual(self.review("--all","--approve").returncode,0)
        self.assertEqual((self.deck/"slide_jobs.json").read_bytes(),current)

    def test_image_drift_invalidates_review_and_assembly(self):
        self.assertEqual(self.review("--all","--approve").returncode,0)
        fixture_image(self.deck/"origin_image/slide_01.png","red")
        self.assertNotEqual(self.assemble().returncode,0)
        self.assertFalse((self.deck/"deck.pptx").exists())

    def test_missing_extra_and_duplicate_final_images_fail_before_output(self):
        self.assertEqual(self.review("--all","--approve").returncode,0)
        path=self.deck/"origin_image/slide_01.png"
        original=path.read_bytes()
        path.unlink()
        self.assertNotEqual(self.assemble().returncode,0)
        path.write_bytes(original)
        for name in ("slide_99.png","slide_1.png","slide_01.jpg"):
            extra=self.deck/"origin_image"/name
            extra.write_bytes(original)
            self.assertNotEqual(self.assemble().returncode,0)
            extra.unlink()
        self.assertFalse((self.deck/"deck.pptx").exists())

    def test_talk_drift_duplicate_heading_and_missing_heading_fail(self):
        self.assertEqual(self.review("--all","--approve").returncode,0)
        path=self.deck/"speech.md"
        original=path.read_text(encoding="utf-8")
        variants=[original+"\nUnapproved closing claim.",
                  original+"\n## Slide 1: Duplicate\n\nHi.",
                  original.split("## Slide 2")[0]]
        for text in variants:
            path.write_text(text,encoding="utf-8")
            self.assertNotEqual(self.assemble().returncode,0)
            self.assertFalse((self.deck/"deck.pptx").exists())

    def test_policy_drift_needs_its_own_actual_approval(self):
        self.assertEqual(self.review("--all","--approve").returncode,0)
        original=build_manifest(self.spec,self.deck)
        revised=copy.deepcopy(self.spec)
        revised["planning"]["brief"].update(output_profile="medavis",landing_title="Shared account")
        write_spec(self.deck/"deck_spec.json",revised)
        self.assertEqual(original["slides"],build_manifest(revised,self.deck)["slides"])
        self.assertNotEqual(self.assemble("--template","medavis","--landing-title","Shared account").returncode,0)
        result=run_cli("record_plan_approval.py",self.deck,"--spec",self.deck/"deck_spec.json",
                       "--approve-policy","--expected-plan-hash",build_manifest(revised,self.deck)["plan_hash"],
                       "--evidence","SYNTHETIC TEST branded output decision")
        self.assertEqual(result.returncode,0,result.stderr)
        result=self.assemble("--template","medavis","--landing-title","Shared account")
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)
        self.assertEqual(len(Presentation(self.deck/"deck-medavis.pptx").slides),4)

    def test_output_ratio_title_profile_mismatches_and_legacy_bypass_fail(self):
        self.assertEqual(self.review("--all","--approve").returncode,0)
        for args in (("--aspect-ratio","4:3"),
                     ("--template","medavis","--landing-title","Changed"),
                     ("--landing-title","Changed"),
                     ("--legacy-unmanaged",)):
            self.assertNotEqual(self.assemble(*args).returncode,0)
            self.assertFalse((self.deck/"deck.pptx").exists())

    def test_complete_sequence_qa_stales_after_approved_talk_revision(self):
        self.assertEqual(self.review("--all","--approve").returncode,0)
        revised=copy.deepcopy(self.spec)
        revised["slides"][1]["speaker_notes"]["text"]+="\n\nLet us discuss."
        approve_fixture(self.deck,revised)
        self.assertNotEqual(self.assemble().returncode,0)
        result=run_cli("record_slide_qa.py",self.deck,"--deck-pass","--evidence","SYNTHETIC TEST revised talk sequence")
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertEqual(self.assemble().returncode,0)

    def test_actual_plain_and_branded_structure_notes_and_uniform_scale(self):
        self.assertEqual(self.review("--all","--approve").returncode,0)
        revised=copy.deepcopy(self.spec)
        revised["planning"]["brief"].update(output_profile="both",landing_title="Shared account")
        approve_fixture(self.deck,revised)
        result=self.assemble("--template","both","--landing-title","Shared account")
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)
        for name,offset,count in (("deck.pptx",0,2),("deck-medavis.pptx",1,4)):
            path=self.deck/name
            prs=Presentation(path)
            self.assertEqual(len(prs.slides),count)
            for i,s in enumerate(self.spec["slides"]):
                slide=prs.slides[i+offset]
                self.assertEqual(slide.notes_slide.notes_text_frame.text,s["speaker_notes"]["text"])
                picture=next(shape for shape in slide.shapes if shape.shape_type==13)
                self.assertAlmostEqual(picture.width/picture.height,16/9,places=4)
            with zipfile.ZipFile(path) as archive:
                self.assertTrue(any(p.startswith("ppt/notesSlides/notesSlide") and p.endswith(".xml") for p in archive.namelist()))
            if offset:
                self.assertNotIn(self.spec["slides"][0]["speaker_notes"]["text"],
                                 prs.slides[0].notes_slide.notes_text_frame.text)
                self.assertNotIn(self.spec["slides"][-1]["speaker_notes"]["text"],
                                 prs.slides[-1].notes_slide.notes_text_frame.text)

    def test_source_changed_during_assembly_preserves_previous_output(self):
        import assemble_ppt
        import io
        from contextlib import redirect_stdout
        self.assertEqual(self.review("--all","--approve").returncode,0)
        output=self.deck/"deck.pptx"
        output.write_bytes(b"Previous approved output")
        original=assemble_ppt.create_presentation
        def changing_source(*args,**kwargs):
            fixture_image(self.deck/"origin_image/slide_01.png","red")
            return original(*args,**kwargs)
        with patch.object(assemble_ppt,"create_presentation",side_effect=changing_source), \
             patch.object(sys,"argv",["assemble_ppt.py",str(self.base),"deck.pptx"]), \
             redirect_stdout(io.StringIO()):
            with self.assertRaises(SystemExit) as result:
                assemble_ppt.main()
        self.assertNotEqual(result.exception.code,0)
        self.assertEqual(output.read_bytes(),b"Previous approved output")

    def test_bulk_review_copy_failure_rolls_back_jobs_and_ledger(self):
        import argparse
        import record_plan_approval
        from record_user_review import record_managed_review
        before={n:(self.deck/n).read_bytes() for n in ("slide_jobs.json","plan_approvals.json")}
        original=record_plan_approval.atomic_bytes
        failed=[]
        def failing_once(path,body):
            if path.name=="slide_jobs.json" and not failed:
                failed.append(True)
                raise OSError("SYNTHETIC TEST state-write failure")
            return original(path,body)
        args=argparse.Namespace(all=True,slides=None,slide=None,exceptions=None,
            approve=True,clear=False,evidence="SYNTHETIC TEST image decision")
        with patch.object(record_plan_approval,"atomic_bytes",side_effect=failing_once):
            with self.assertRaises(OSError):
                record_managed_review(self.deck,args)
        self.assertEqual(before,{n:(self.deck/n).read_bytes() for n in before})
        self.assertFalse((self.deck/".planning_transaction.json").exists())
