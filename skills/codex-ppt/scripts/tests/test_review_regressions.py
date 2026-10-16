"""Independent review regressions; all decisions and pixels are synthetic."""
import copy
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from unittest.mock import patch

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))
from planning_fixtures import (approve_fixture, compile_fixture, dispatch_fixture,
    fixture_image, managed_spec, qa_fixture, result_fixture, run_cli, write_spec)
import prepare_slide_prompts
import record_plan_approval
import record_slide_result
import slide_run_state


class ReviewRegressions(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.deck = Path(self.tmp.name).resolve() / "deck"
        self.deck.mkdir()

    def checked(self, result):
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return result

    def compile(self, phase="production", slides="slide_02", **kwargs):
        args = ["--spec", self.deck / "deck_spec.json", "--out-dir", self.deck,
                "--phase", phase, "--slides", slides, "--force"]
        if "capacity" in kwargs:
            args += ["--max-concurrent-slides", kwargs["capacity"]]
        return run_cli("prepare_slide_prompts.py", *args)

    def record(self, sid="slide_01", color="white"):
        self.checked(dispatch_fixture(self.deck, sid))
        self.checked(result_fixture(self.deck, sid, color=color))

    def jobs(self):
        return json.loads((self.deck / "slide_jobs.json").read_text(encoding="utf-8"))

    def test_accepted_sample_repair_dispatches_and_preserves_style_snapshot(self):
        compile_fixture(self.deck)
        ledger = json.loads((self.deck / "plan_approvals.json").read_text())
        reference = self.deck / ledger["samples"]["slide_02"]["image"]
        before = reference.read_bytes()
        self.checked(qa_fixture(self.deck, "slide_02", "--needs-repair"))
        self.checked(self.compile("repair"))
        self.assertEqual(self.jobs()["slides"][1]["status"], "pending")
        self.record("slide_02", "red")
        self.checked(qa_fixture(self.deck, "slide_02"))
        self.assertEqual(reference.read_bytes(), before)
        self.checked(run_cli("record_user_review.py", self.deck, "--slide", "slide_02",
                             "--approve", "--evidence", "SYNTHETIC repaired final decision"))

    def test_unapproved_calibration_repairs_without_approving_defective_sample(self):
        approve_fixture(self.deck)
        self.checked(self.compile("sample"))
        self.record("slide_02")
        self.checked(qa_fixture(self.deck, "slide_02", "--needs-repair"))
        self.checked(self.compile("repair"))
        self.record("slide_02", "red")
        self.checked(qa_fixture(self.deck, "slide_02"))
        self.assertEqual(self.jobs()["slides"][1]["repair_counts"].values().__iter__().__next__(), 1)
        ledger = json.loads((self.deck / "plan_approvals.json").read_text())
        self.assertEqual(ledger["samples"], {})
        self.checked(run_cli("record_plan_approval.py", self.deck, "--approve-sample", "slide_02",
            "--image", self.deck / "origin_image/slide_02.png", "--backend-used", "scripts/image_gen.py",
            "--evidence", "SYNTHETIC repaired calibration decision"))
        self.checked(run_cli("prepare_slide_prompts.py", "--spec", self.deck / "deck_spec.json",
                             "--out-dir", self.deck, "--force"))
        self.assertEqual(self.jobs()["slides"][1]["phase"], "repair")
        self.record(); self.checked(qa_fixture(self.deck))
        self.checked(run_cli("record_slide_qa.py", self.deck, "--deck-pass",
                             "--evidence", "SYNTHETIC repaired calibration full sequence"))
        self.checked(run_cli("record_user_review.py", self.deck, "--all", "--approve",
                             "--evidence", "SYNTHETIC whole current deck decision"))
        self.checked(run_cli("assemble_ppt.py", self.deck.parent, "deck.pptx"))

    def test_missing_tracked_prompt_cannot_bypass_repair_or_dispatch(self):
        compile_fixture(self.deck)
        self.record()
        self.checked(qa_fixture(self.deck, decision="--needs-repair"))
        before = (self.deck / "slide_jobs.json").read_bytes()
        (self.deck / "prompts/slide_01.json").unlink()
        result = self.compile("production", "slide_01")
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual((self.deck / "slide_jobs.json").read_bytes(), before)

    def test_first_batch_revision_regenerates_before_later_production(self):
        spec = managed_spec()
        third = copy.deepcopy(spec["slides"][0]); third["number"] = 3
        spec["slides"].append(third)
        compile_fixture(self.deck, spec)
        self.record(); self.checked(qa_fixture(self.deck))
        spec["slides"][0]["visible_copy"]["body"] = ["Approved revised phrase"]
        approve_fixture(self.deck, spec)
        self.checked(self.compile("production", "slide_01"))
        self.assertNotEqual(dispatch_fixture(self.deck, "slide_03").returncode, 0)
        self.record()
        self.assertNotEqual(dispatch_fixture(self.deck, "slide_03").returncode, 0)
        self.checked(qa_fixture(self.deck))
        self.checked(dispatch_fixture(self.deck, "slide_03"))

    def test_assembly_uses_plan_order_after_job_array_reordering(self):
        compile_fixture(self.deck)
        self.record(color="red")
        for sid in ("slide_01", "slide_02"):
            self.checked(qa_fixture(self.deck, sid))
        self.checked(run_cli("record_slide_qa.py", self.deck, "--deck-pass",
                             "--evidence", "SYNTHETIC sequence inspection"))
        self.checked(run_cli("record_user_review.py", self.deck, "--all", "--approve",
                             "--evidence", "SYNTHETIC image decision"))
        jobs = self.jobs(); jobs["slides"].reverse()
        write_spec(self.deck / "slide_jobs.json", jobs)
        self.checked(run_cli("assemble_ppt.py", self.deck.parent, "deck.pptx"))
        from pptx import Presentation
        from PIL import Image
        prs = Presentation(self.deck / "deck.pptx")
        image = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
        self.assertEqual(Image.open(io.BytesIO(image.image.blob)).getpixel((0, 0)), (255, 0, 0))
        self.assertIn("Hi everyone.", prs.slides[0].notes_slide.notes_text_frame.text)

    def test_partial_result_copy_preserves_previous_image_and_jobs(self):
        compile_fixture(self.deck)
        target = self.deck / "origin_image/slide_01.png"
        fixture_image(target, "blue")
        self.checked(dispatch_fixture(self.deck))
        source = self.deck / "result.png"; fixture_image(source, "red")
        from planning_fixtures import register_fixture_receipt
        receipt=register_fixture_receipt(self.deck,"slide_01",source)
        before = (target.read_bytes(), (self.deck / "slide_jobs.json").read_bytes())
        def broken_copy(source, destination):
            Path(destination).write_bytes(b"incomplete copy")
            raise OSError("SYNTHETIC partial copy failure")
        with patch.object(sys, "argv", ["record_slide_result.py", str(self.deck), "--slide", "1",
            "--agent-id", "fixture-agent", "--backend-used", "scripts/image_gen.py",
            "--selected-source", str(source), "--submission-receipt", str(receipt), "--qa-note", "SYNTHETIC"]), \
            patch.object(record_slide_result.shutil, "copy2", side_effect=broken_copy), redirect_stdout(io.StringIO()):
            with self.assertRaises(OSError):
                record_slide_result.main()
        self.assertEqual((target.read_bytes(), (self.deck / "slide_jobs.json").read_bytes()), before)

    def test_compilation_rolls_back_prompt_and_jobs_write_failures(self):
        for failed_name in ("slide_01.json", "slide_02.json", "slide_jobs.json"):
            with self.subTest(failed_name=failed_name):
                self.deck = Path(self.tmp.name).resolve() / failed_name
                self.deck.mkdir()
                spec = compile_fixture(self.deck)
                before = {p: p.read_bytes() for p in [self.deck / "slide_jobs.json",
                    self.deck / "prompts/slide_01.json", self.deck / "prompts/slide_02.json"]}
                spec["style"]["color_palette"] = "approved orange and white"
                approve_fixture(self.deck, spec)
                original = record_plan_approval.atomic_bytes
                failed = []
                def failing_write(path, body):
                    if path.name == failed_name and not failed:
                        failed.append(True)
                        raise OSError("SYNTHETIC grouped write failure")
                    return original(path, body)
                # Reapprove a matching synthetic sample, then compile both changed jobs.
                from planning_fixtures import approve_fixture_sample
                approve_fixture_sample(self.deck)
                # Change approved page content after calibration so both packets
                # must be rewritten, while its immutable style sample stays valid.
                for slide in spec["slides"]:
                    slide["title"] += " revised"
                    slide["visible_copy"]["title"] = slide["title"]
                approve_fixture(self.deck,spec)
                before = {p:p.read_bytes() for p in before}
                with patch.object(sys, "argv", ["prepare_slide_prompts.py", "--spec",
                    str(self.deck / "deck_spec.json"), "--out-dir", str(self.deck), "--force"]), \
                    patch.object(record_plan_approval, "atomic_bytes", side_effect=failing_write), redirect_stdout(io.StringIO()):
                    with self.assertRaises(OSError):
                        prepare_slide_prompts.main()
                for path, body in before.items():
                    self.assertEqual(path.read_bytes(), body)

    def test_result_state_write_failure_preserves_previous_image_and_state(self):
        compile_fixture(self.deck)
        target=self.deck/"origin_image/slide_01.png";fixture_image(target,"blue")
        self.checked(dispatch_fixture(self.deck))
        source=self.deck/"new-result.png";fixture_image(source,"red")
        from planning_fixtures import register_fixture_receipt
        receipt=register_fixture_receipt(self.deck,"slide_01",source)
        before=(target.read_bytes(),(self.deck/"slide_jobs.json").read_bytes())
        original=record_plan_approval.atomic_bytes;failed=[]
        def fail_state_once(path,body):
            if path.name=="slide_jobs.json" and not failed:
                failed.append(True);raise OSError("SYNTHETIC result state write failure")
            return original(path,body)
        with patch.object(sys,"argv",["record_slide_result.py",str(self.deck),"--slide","1",
            "--agent-id","fixture-agent","--backend-used","scripts/image_gen.py",
            "--selected-source",str(source),"--submission-receipt",str(receipt),"--qa-note","SYNTHETIC"]), \
            patch.object(record_plan_approval,"atomic_bytes",side_effect=fail_state_once),redirect_stdout(io.StringIO()):
            with self.assertRaises(OSError):
                record_slide_result.main()
        self.assertEqual((target.read_bytes(),(self.deck/"slide_jobs.json").read_bytes()),before)
        self.assertFalse((self.deck/".planning_transaction.json").exists())

    def test_zero_capacity_is_rejected_without_state_mutation(self):
        compile_fixture(self.deck)
        before = (self.deck / "slide_jobs.json").read_bytes()
        self.assertNotEqual(self.compile(capacity=0).returncode, 0)
        self.assertEqual((self.deck / "slide_jobs.json").read_bytes(), before)

    def test_compiler_and_blocker_do_not_lose_interleaved_state(self):
        compile_fixture(self.deck)
        original = slide_run_state.read_json
        children = []
        def read_and_start_blocker(path, *args, **kwargs):
            value = original(path, *args, **kwargs)
            if path.name == "slide_jobs.json" and not children:
                process = subprocess.Popen([sys.executable, "-B", str(SCRIPTS / "record_slide_blocker.py"),
                    str(self.deck), "--slide", "1", "--reason", "SYNTHETIC unavailable backend"],
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8",
                    env=dict(os.environ, PYTHONIOENCODING="utf-8"))
                children.append(process)
                try:
                    process.communicate(timeout=1)
                except subprocess.TimeoutExpired:
                    pass  # Correct locking holds the worker until compiler releases its locks.
            return value
        with patch.object(sys, "argv", ["prepare_slide_prompts.py", "--spec",
            str(self.deck / "deck_spec.json"), "--out-dir", str(self.deck), "--max-concurrent-slides", "2"]), \
            patch.object(slide_run_state, "read_json", side_effect=read_and_start_blocker), redirect_stdout(io.StringIO()):
            prepare_slide_prompts.main()
        output, error = children[0].communicate(timeout=10)
        self.assertEqual(children[0].returncode, 0, error)
        self.assertEqual(self.jobs()["slides"][0]["status"], "blocked")
        self.assertIn("unavailable backend", self.jobs()["slides"][0]["blocker"]["reason"])
